"""Generate PDF and editable DOCX files from one frozen assessment snapshot."""

from html import escape
from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Image, PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

FIELD_LABELS = {
    "product_name": "Product name",
    "manufacturer": "Manufacturer / packer / importer",
    "address": "Responsible business address",
    "net_quantity": "Net quantity",
    "mrp": "Maximum retail price",
    "date": "Manufacture / packing / import date",
    "consumer_care": "Consumer care",
    "country_of_origin": "Country of origin",
}


def _text(value) -> str:
    return str(value if value not in (None, "") else "Not recorded")


def _pdf_text(value) -> str:
    """Escape dynamic text before ReportLab parses its small markup language."""
    return escape(_text(value))


def _image_path(data_dir: Path, item: dict) -> Path | None:
    name = Path(str(item.get("stored_name", ""))).name
    candidate = data_dir / "images" / name
    return candidate if name and candidate.is_file() else None


def _evidence_label(finding: dict) -> str:
    evidence = finding.get("evidence") or {}
    if not evidence:
        return "No direct image evidence linked"
    return f"{_text(evidence.get('panel')).title()} panel / {_text(evidence.get('filename'))} / image {_text(evidence.get('image_id'))[:8]}"


def build_pdf(snapshot: dict, output: Path, data_dir: Path) -> Path:
    styles = getSampleStyleSheet()
    styles.add(
        ParagraphStyle(
            name="ReportTitle",
            parent=styles["Title"],
            textColor=colors.HexColor("#133C3D"),
            fontSize=22,
            leading=27,
            alignment=TA_CENTER,
        )
    )
    styles.add(
        ParagraphStyle(
            name="Section",
            parent=styles["Heading2"],
            textColor=colors.HexColor("#133C3D"),
            fontSize=13,
            spaceBefore=10,
            spaceAfter=6,
        )
    )
    styles.add(
        ParagraphStyle(
            name="Small", parent=styles["BodyText"], fontSize=8, leading=10, textColor=colors.HexColor("#536466")
        )
    )
    styles.add(ParagraphStyle(name="TableCell", parent=styles["BodyText"], fontSize=8, leading=10))
    styles.add(
        ParagraphStyle(name="TableHeader", parent=styles["TableCell"], textColor=colors.white, fontName="Helvetica-Bold")
    )
    document = SimpleDocTemplate(
        str(output),
        pagesize=A4,
        leftMargin=18 * mm,
        rightMargin=18 * mm,
        topMargin=16 * mm,
        bottomMargin=16 * mm,
        title=f"CompliSense assessment - {_text(snapshot.get('product_name'))}",
    )
    story = [
        Paragraph("CompliSense assessment report", styles["ReportTitle"]),
        Paragraph(
            "Evidence-linked prototype assessment. This report is not a legal compliance certificate.", styles["Small"]
        ),
        Spacer(1, 7 * mm),
    ]
    summary = [
        ["Product", _text(snapshot.get("product_name"))],
        ["Brand", _text(snapshot.get("brand"))],
        ["Inspection", _text(snapshot.get("id"))],
        ["Inspector", _text(snapshot.get("owner_name"))],
        ["Rule set", _text(snapshot.get("rule_version"))],
        ["Assessment version", _text(snapshot.get("assessment_version") or snapshot.get("assessment_id"))],
        ["Panel coverage confirmed", "Yes" if snapshot.get("coverage_confirmed") else "No"],
        ["Applicability", _text((snapshot.get("applicability") or {}).get("status"))],
        ["Exemption record", _text((snapshot.get("applicability") or {}).get("exemption"))],
        ["Review decision", _text(snapshot.get("review_decision"))],
    ]
    table = Table(summary, colWidths=[48 * mm, 112 * mm])
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#E8F0EC")),
                ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#B7C4C0")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("FONTNAME", (0, 0), (0, -1), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, -1), 9),
                ("PADDING", (0, 0), (-1, -1), 6),
            ]
        )
    )
    story += [table, Paragraph("Extracted declarations", styles["Section"])]
    declaration_rows = [
        [
            Paragraph("Declaration", styles["TableHeader"]),
            Paragraph("Recorded value", styles["TableHeader"]),
            Paragraph("Source", styles["TableHeader"]),
        ]
    ]
    for key, declaration in snapshot.get("fields", {}).items():
        declaration_rows.append(
            [
                Paragraph(_pdf_text(FIELD_LABELS.get(key, key.replace("_", " ").title())), styles["TableCell"]),
                Paragraph(_pdf_text(declaration.get("value")), styles["TableCell"]),
                Paragraph(_pdf_text(declaration.get("source")), styles["TableCell"]),
            ]
        )
    declarations = Table(declaration_rows, repeatRows=1, colWidths=[45 * mm, 93 * mm, 22 * mm])
    declarations.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#133C3D")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#BBC6C3")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("FONTSIZE", (0, 0), (-1, -1), 8),
                ("PADDING", (0, 0), (-1, -1), 5),
            ]
        )
    )
    story += [declarations, Paragraph("Findings", styles["Section"])]
    for finding in snapshot.get("findings", []):
        story += [
            Paragraph(
                f"<b>{_pdf_text(finding.get('rule_id'))} - {_pdf_text(finding.get('title'))}</b> - {_pdf_text(finding.get('status')).replace('_', ' ').title()}",
                styles["BodyText"],
            ),
            Paragraph(_pdf_text(finding.get("explanation")), styles["Small"]),
            Paragraph(
                f"Evidence: {_pdf_text(_evidence_label(finding))}",
                styles["Small"],
            ),
            Paragraph(
                f"Rule reference: {_pdf_text(finding.get('provision'))} - {_pdf_text(finding.get('source_url'))}",
                styles["Small"],
            ),
            Spacer(1, 3 * mm),
        ]
    story += [Paragraph("Manual checks and review", styles["Section"])]
    for key, value in snapshot.get("manual_checks", {}).items():
        story.append(
            Paragraph(f"<b>{escape(key.replace('_', ' ').title())}:</b> {_pdf_text(value)}", styles["BodyText"])
        )
    story.append(Paragraph(f"<b>Reviewer notes:</b> {_pdf_text(snapshot.get('review_notes'))}", styles["BodyText"]))
    for item in snapshot.get("images", []):
        path = _image_path(data_dir, item)
        if path:
            story += [
                PageBreak(),
                Paragraph(f"Evidence: {_pdf_text(item.get('panel')).title()} panel", styles["Section"]),
                Image(str(path), width=160 * mm, height=105 * mm, kind="proportional"),
                Paragraph(f"File hash (SHA-256): {_pdf_text(item.get('sha256'))}", styles["Small"]),
            ]
    document.build(story)
    return output


def build_docx(snapshot: dict, output: Path, data_dir: Path) -> Path:
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    document.styles["Normal"].font.name = "Aptos"
    document.styles["Normal"].font.size = Pt(10)
    title = document.add_heading("CompliSense assessment report", 0)
    title.alignment = 1
    document.add_paragraph(
        "Evidence-linked prototype assessment. This editable report is not a legal compliance certificate."
    )
    table = document.add_table(rows=0, cols=2)
    table.style = "Light Shading Accent 1"
    for label, value in (
        ("Product", snapshot.get("product_name")),
        ("Brand", snapshot.get("brand")),
        ("Inspection", snapshot.get("id")),
        ("Inspector", snapshot.get("owner_name")),
        ("Rule set", snapshot.get("rule_version")),
        ("Assessment version", snapshot.get("assessment_version") or snapshot.get("assessment_id")),
        ("Panel coverage confirmed", "Yes" if snapshot.get("coverage_confirmed") else "No"),
        ("Applicability", (snapshot.get("applicability") or {}).get("status")),
        ("Exemption record", (snapshot.get("applicability") or {}).get("exemption")),
        ("Review decision", snapshot.get("review_decision")),
    ):
        cells = table.add_row().cells
        cells[0].text = label
        cells[1].text = _text(value)
    document.add_heading("Extracted declarations", level=1)
    declarations = document.add_table(rows=1, cols=3)
    declarations.style = "Light Grid Accent 1"
    for cell, value in zip(
        declarations.rows[0].cells,
        ("Declaration", "Recorded value", "Source"),
        strict=True,
    ):
        cell.text = value
    for key, declaration in snapshot.get("fields", {}).items():
        cells = declarations.add_row().cells
        cells[0].text = FIELD_LABELS.get(key, key.replace("_", " ").title())
        cells[1].text = _text(declaration.get("value"))
        cells[2].text = _text(declaration.get("source"))
    document.add_heading("Findings", level=1)
    for finding in snapshot.get("findings", []):
        document.add_heading(f"{_text(finding.get('rule_id'))} - {_text(finding.get('title'))}", level=2)
        document.add_paragraph(f"Status: {_text(finding.get('status')).replace('_', ' ').title()}")
        document.add_paragraph(_text(finding.get("explanation")))
        document.add_paragraph(f"Evidence: {_evidence_label(finding)}")
        document.add_paragraph(f"Rule reference: {_text(finding.get('provision'))} - {_text(finding.get('source_url'))}")
    document.add_heading("Manual checks and review", level=1)
    for key, value in snapshot.get("manual_checks", {}).items():
        document.add_paragraph(f"{key.replace('_', ' ').title()}: {_text(value)}")
    document.add_paragraph(f"Reviewer notes: {_text(snapshot.get('review_notes'))}")
    for item in snapshot.get("images", []):
        path = _image_path(data_dir, item)
        if path:
            document.add_page_break()
            document.add_heading(f"Evidence: {_text(item.get('panel')).title()} panel", level=1)
            document.add_picture(str(path), width=Inches(6.5))
            document.add_paragraph(f"File hash (SHA-256): {_text(item.get('sha256'))}")
    document.save(output)
    return output

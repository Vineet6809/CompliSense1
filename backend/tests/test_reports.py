"""Both report formats must contain the same frozen assessment facts."""

from io import BytesIO

from app.services.reports import build_docx, build_pdf
from docx import Document
from pypdf import PdfReader

SNAPSHOT = {
    "id": "inspection-1",
    "product_name": "Example Soap",
    "brand": "Demo Brand",
    "owner_name": "Inspector One",
    "created_at": "2026-09-17T10:00:00+00:00",
    "rule_version": "demo-2026.09",
    "assessment_version": "a1b2c3d4",
    "applicability": {"status": "applicable", "exemption": "none", "reason": "Household commodity scope verified."},
    "coverage_confirmed": True,
    "review_decision": "follow_up",
    "review_notes": "Confirm address against original pack.",
    "fields": {"net_quantity": {"value": "100 g", "confidence": 0.94, "source": "ocr", "note": ""}},
    "findings": [
        {
            "rule_id": "LMPC-6-1-C",
            "title": "Net quantity declaration",
            "status": "pass",
            "explanation": "A candidate declaration was extracted.",
            "provision": "Rule 6(1)(c)",
            "source_url": "https://consumeraffairs.gov.in/pages/legal-metrology-act",
            "evidence": {"image_id": "image-1", "filename": "front-label.png", "panel": "front", "sha256": "abc123"},
        }
    ],
    "manual_checks": {"font_size": "unassessed", "placement": "unassessed", "readability": "acceptable", "notes": ""},
    "images": [],
    "processing_seconds": 1.2,
}


def test_pdf_contains_product_rule_and_review_note(tmp_path):
    output = tmp_path / "report.pdf"
    build_pdf(SNAPSHOT, output, tmp_path)
    text = "\n".join(page.extract_text() or "" for page in PdfReader(output).pages)
    assert "Example Soap" in text
    assert "LMPC-6-1-C - Net quantity declaration - Pass" in text
    assert "Confirm address" in text
    assert "front-label.png" in text
    assert "a1b2c3d4" in text
    assert "�" not in text


def test_docx_contains_product_rule_and_review_note(tmp_path):
    output = tmp_path / "report.docx"
    build_docx(SNAPSHOT, output, tmp_path)
    document = Document(BytesIO(output.read_bytes()))
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    text += "\n" + "\n".join(cell.text for table in document.tables for row in table.rows for cell in row.cells)
    assert "Example Soap" in text
    assert "LMPC-6-1-C - Net quantity declaration" in text
    assert "Confirm address" in text
    assert "front-label.png" in text
    assert "a1b2c3d4" in text
